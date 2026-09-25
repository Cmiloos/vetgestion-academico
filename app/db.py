"""Datos de VetGestión: productos, stock y movimientos.

La información vive en uno de estos dos lugares ("almacenes"):
  - Google Sheets (en la nube, compartido entre computadoras) si está configurado.
  - Excel local datos/vetgestion.xlsx (respaldo sin internet), si no.

Las dos planillas tienen las mismas hojas:
  Productos    una fila por producto (datos, precio, stock, vencimiento, ficha)
  Movimientos  historial: ventas, aumentos, altas, ediciones, reposiciones, bajas
"""
import base64
import copy
import json
import os
import shutil
import sys
import threading
import time
import urllib.error
import urllib.request
from datetime import datetime, timedelta

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

# Columnas de la hoja Productos: (clave interna, encabezado visible, ancho)
COLS_PROD = [
    ("codigo", "Código", 11),
    ("nombre", "Producto", 24),
    ("laboratorio", "Laboratorio", 14),
    ("categoria", "Categoría", 22),
    ("especie", "Especie", 13),
    ("presentacion", "Presentación", 13),
    ("droga", "Droga", 22),
    ("precio", "Precio", 11),
    ("stock", "Stock", 8),
    ("vencimiento", "Vencimiento (MM/AAAA)", 12),
    ("concentracion", "Concentración", 26),
    ("composicion", "Composición", 50),
    ("via", "Vía", 16),
    ("dosis", "Dosis sugerida", 26),
    ("lote", "Lote", 11),
    ("indicaciones", "Indicaciones", 50),
    ("conservacion", "Conservación", 36),
    ("imagen", "Imagen (archivo)", 22),
]
COLS_MOV = [
    ("fecha", "Fecha", 18),
    ("tipo", "Tipo", 20),
    ("codigo", "Código", 11),
    ("producto", "Producto", 26),
    ("cantidad", "Cantidad", 10),
    ("importe", "Importe", 12),
    ("detalle", "Detalle", 44),
    ("usuario", "Usuario", 12),
]
NUMERICAS = {"precio", "stock", "cantidad", "importe"}
CON_PESOS = {"precio", "importe"}

STOCK_BAJO = 5          # stock igual o menor a esto = "stock bajo"
DIAS_PRONTO = 90        # vence dentro de estos días = "pronto"
CACHE_SEG = 3           # una lectura sirve este tiempo (evita pedir dos veces seguidas a Sheets)
FMT_FECHA = "%Y-%m-%d %H:%M:%S"

_lock = threading.Lock()


def _base_dir():
    """Carpeta donde vive el programa (junto al .exe, o la raíz del proyecto)."""
    if getattr(sys, "frozen", False):
        return os.path.dirname(sys.executable)
    return os.path.dirname(os.path.dirname(os.path.abspath(__file__)))


def _recursos_dir():
    """Carpeta con los archivos empaquetados (ui/, imágenes de ejemplo)."""
    return getattr(sys, "_MEIPASS", os.path.dirname(os.path.abspath(__file__)))


def _datos_dir():
    """El .exe guarda sus datos en la carpeta de la app de Windows, así no crea nada al lado del programa
    (se puede tener suelto en el Escritorio o en un pendrive). En desarrollo se usa la carpeta datos/."""
    if getattr(sys, "frozen", False):
        return os.path.join(os.environ.get("LOCALAPPDATA") or os.path.expanduser("~"), "VetGestion")
    return os.path.join(_base_dir(), "datos")


DATOS_DIR = _datos_dir()
IMG_DIR = os.path.join(DATOS_DIR, "imagenes")
XLSX = os.path.join(DATOS_DIR, "vetgestion.xlsx")
CONFIG = os.path.join(DATOS_DIR, "config.json")


class ExcelAbierto(ValueError):
    pass


class SinConexion(ValueError):
    pass


# ---------------------------------------------------------------- datos de ejemplo
SEED = [
    ("005-0142", "Osteovet comprimidos", "AndinaVet", "Óseo / articular", "Perro", "Comprimido", "Carprofeno", 8900, 0, "08/2026",
     "75 mg por comprimido", "Carprofeno 75 mg · Lactosa monohidrato · Celulosa microcristalina · Estearato de magnesio", "Oral", "2 mg/kg cada 12 h", "L-2405-A",
     "Antiinflamatorio y analgésico para procesos osteoarticulares crónicos en caninos.", "Lugar seco, por debajo de 25 °C.", "005-0142.webp"),
    ("005-0210", "Cardiovet 5 mg", "AndinaVet", "Cardíaco", "Perro", "Comprimido", "Pimobendan", 15400, 12, "07/2027",
     "5 mg por comprimido", "Pimobendan 5 mg · Almidón de maíz · Povidona · Estearato de magnesio", "Oral", "0,25 mg/kg cada 12 h", "L-2602-C",
     "Insuficiencia cardíaca congestiva por cardiomiopatía dilatada o insuficiencia valvular.", "Lugar seco, por debajo de 25 °C.", "005-0210.webp"),
    ("006-0301", "Pulgex pipeta", "ZooFarma", "Antiparasitario externo", "Gato", "Pipeta", "Fipronil", 6200, 8, "10/2026",
     "0,67 ml · Fipronil 10 %", "Fipronil 67 mg · Butilhidroxianisol · Etanol", "Tópica (spot-on)", "1 pipeta por mes", "L-2511-P",
     "Pulgas y garrapatas en felinos de hasta 4 kg.", "Lugar fresco, lejos de la luz directa.", "006-0301.webp"),
    ("006-0088", "Antipar Plus", "ZooFarma", "Antiparasitario externo", "Perro", "Pipeta", "Fipronil + permetrina", 4700, 3, "01/2027",
     "1,1 ml · Fipronil 10 % + permetrina 45 %", "Fipronil 110 mg · Permetrina 495 mg · Excipientes c.s.p. 1,1 ml", "Tópica (spot-on)", "1 pipeta por mes", "L-2512-N",
     "Pulgas, garrapatas y mosquitos en caninos de hasta 10 kg. No usar en gatos.", "Lugar fresco, lejos de la luz directa.", "006-0088.jpg"),
    ("007-0155", "Meloxivet 15 ml", "PetLab", "Antiinflamatorio", "Perro y gato", "Jarabe", "Meloxicam", 9800, 21, "05/2027",
     "1,5 mg/ml · frasco de 15 ml", "Meloxicam 1,5 mg/ml · Sorbitol · Glicerol · Benzoato de sodio", "Oral", "0,1 mg/kg una vez al día", "L-2604-M",
     "Antiinflamatorio no esteroide para dolor agudo y crónico en perros y gatos.", "Agitar antes de usar. Consumir dentro de los 6 meses de abierto.", "007-0155.jpg"),
    ("007-0402", "Amoxivet 500", "PetLab", "Antibiótico", "Perro", "Comprimido", "Amoxicilina", 7300, 2, "11/2027",
     "500 mg por cápsula", "Amoxicilina trihidrato 500 mg · Dióxido de silicio · Gelatina", "Oral", "10 mg/kg cada 12 h", "L-2603-A",
     "Infecciones bacterianas de piel, vías urinarias y respiratorias.", "Lugar seco, por debajo de 25 °C.", "007-0402.jpg"),
    ("008-0019", "Dermivet loción", "CanVet", "Dermatológico", "Perro y gato", "Loción", "Clorhexidina", 5600, 0, "07/2026",
     "2 g/100 ml · frasco de 100 ml", "Gluconato de clorhexidina 2 g · Glicerina · Alantoína · Agua purificada", "Tópica", "Aplicar 1 a 2 veces al día", "L-2408-D",
     "Dermatitis bacterianas y seborreicas en perros y gatos.", "No aplicar sobre heridas profundas.", "008-0019.jpg"),
    ("008-0233", "Gastrovet", "CanVet", "Digestivo", "Perro", "Jarabe", "Omeprazol", 6100, 4, "02/2027",
     "4 mg/ml · frasco de 60 ml", "Omeprazol 4 mg/ml · Extracto de regaliz · Sorbitol · Saborizante", "Oral", "1 mg/kg una vez al día, en ayunas", "L-2601-G",
     "Gastritis, reflujo y úlceras gástricas en caninos.", "Conservar refrigerado una vez abierto.", "008-0233.jpg"),
    ("009-0077", "Vermicán jarabe", "NorVet", "Antiparasitario interno", "Gato", "Jarabe", "Pamoato de pirantel", 5200, 16, "09/2027",
     "50 mg/ml · frasco de 20 ml", "Pamoato de pirantel 50 mg/ml · Praziquantel 15 mg/ml · Saborizante dulce de leche", "Oral", "1 ml cada 2 kg, dosis única", "L-2605-V",
     "Áscaris, anquilostomas y tenias en gatos y cachorros.", "Agitar bien antes de administrar.", "009-0077.jpg"),
    ("009-0500", "Anticoncep-D", "NorVet", "Anticonceptivo", "Perro", "Comprimido", "Acetato de megestrol", 8100, 9, "11/2026",
     "5 mg por comprimido", "Acetato de megestrol 5 mg · Lactosa · Talco · Estearato de magnesio", "Oral", "Según protocolo veterinario", "L-2510-N",
     "Control del ciclo estral en caninas. Requiere indicación profesional.", "Lugar seco, fuera del alcance de niños.", "009-0500.jpg"),
    ("010-0044", "Lechevit cachorros", "BioPet", "Nutrición", "Perro", "Polvo", "—", 11200, 18, "04/2027",
     "400 g · rinde 3,2 litros", "Leche entera en polvo · Vitaminas A y D · Calcio · Taurina · Lecitina de soja", "Oral", "1 medida cada 50 ml de agua tibia", "L-2606-L",
     "Sustituto lácteo para cachorros huérfanos o en destete temprano.", "Cerrar bien el envase. Consumir dentro de los 30 días de abierto.", "010-0044.jpg"),
    ("010-0126", "Otivet gotas", "BioPet", "Dermatológico", "Perro y gato", "Gotas", "Ketoconazol", 6900, 6, "06/2027",
     "15 ml · solución ótica", "Ketoconazol 10 mg/ml · Ciprofloxacina 3 mg/ml · Dexametasona 1 mg/ml", "Ótica", "3 a 5 gotas cada 12 h", "L-2602-O",
     "Otitis externa bacteriana y micótica en perros y gatos.", "No usar si hay perforación timpánica.", "010-0126.jpg"),
]


# ---------------------------------------------------------------- normalización
def _claves(cols):
    return [k for k, _, _ in cols]


def _a_fila(d, cols):
    return [d.get(k, "") for k in _claves(cols)]


def _entero(v):
    try:
        return int(round(float(str(v).replace(",", ".")))) if v not in (None, "") else 0
    except (TypeError, ValueError):
        return 0


def _limpiar_producto(p):
    out = {}
    for k in _claves(COLS_PROD):
        v = p.get(k)
        if k in NUMERICAS:
            out[k] = _entero(v)
        elif isinstance(v, datetime):  # fecha escrita a mano en la planilla
            out[k] = v.strftime("%m/%Y") if k == "vencimiento" else v.strftime("%d/%m/%Y")
        else:
            out[k] = "" if v is None else str(v).strip()
    # "2027-05-01 00:00:00" (fecha escrita a mano en Sheets) -> "05/2027"
    v = out["vencimiento"]
    if len(v) >= 7 and v[4] == "-" and v[:4].isdigit():
        out["vencimiento"] = f"{v[5:7]}/{v[:4]}"
    return out


def _limpiar_mov(m):
    out = {}
    for k in _claves(COLS_MOV):
        v = m.get(k)
        if k in NUMERICAS:
            out[k] = _entero(v)
        elif isinstance(v, datetime):
            out[k] = v.strftime(FMT_FECHA)
        else:
            out[k] = "" if v is None else str(v).strip()
    return out


def estado_vencimiento(venc, hoy=None):
    """Devuelve ('vencido'|'pronto'|'ok'|'sin_dato', días hasta el vencimiento)."""
    hoy = hoy or datetime.now().date()
    if isinstance(venc, datetime):
        venc = venc.strftime("%m/%Y")
    try:
        mes, anio = str(venc).split("/")
        mes, anio = int(mes), int(anio)
        # vence el último día del mes indicado
        fin = (datetime(anio + (mes // 12), mes % 12 + 1, 1) - timedelta(days=1)).date()
    except (ValueError, TypeError):
        return "sin_dato", None
    dias = (fin - hoy).days
    if dias < 0:
        return "vencido", dias
    if dias <= DIAS_PRONTO:
        return "pronto", dias
    return "ok", dias


# ---------------------------------------------------------------- almacén: Excel local
def _estilar(ws, cols):
    head_fill = PatternFill("solid", fgColor="123A30")
    for i, (_, titulo, ancho) in enumerate(cols, start=1):
        c = ws.cell(row=1, column=i, value=titulo)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = head_fill
        c.alignment = Alignment(vertical="center")
        ws.column_dimensions[get_column_letter(i)].width = ancho
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}1"


def _formatear_filas(ws, cols):
    idx = {k: i for i, (k, _, _) in enumerate(cols, start=1)}
    for r in range(2, ws.max_row + 1):
        for k in CON_PESOS & idx.keys():
            ws.cell(row=r, column=idx[k]).number_format = '"$"#,##0'
        if "fecha" in idx:
            ws.cell(row=r, column=idx["fecha"]).number_format = "dd/mm/yyyy hh:mm"


def _filas_excel(ws, cols):
    claves = _claves(cols)
    return [dict(zip(claves, row)) for row in ws.iter_rows(min_row=2, max_col=len(cols), values_only=True)
            if row and row[0] not in (None, "")]


def _buscar(ws, codigo):
    for r in range(2, ws.max_row + 1):
        if str(ws.cell(row=r, column=1).value or "").strip() == codigo:
            return r
    return None


class ExcelStore:
    modo = "excel"
    nombre = "Excel"

    def _mov_excel(self, m):
        fila = _a_fila(m, COLS_MOV)
        try:
            fila[0] = datetime.strptime(fila[0], FMT_FECHA)
        except (TypeError, ValueError):
            pass
        return fila

    def _guardar(self, wb):
        for ws, cols in ((wb["Productos"], COLS_PROD), (wb["Movimientos"], COLS_MOV)):
            _formatear_filas(ws, cols)
        try:
            wb.save(XLSX)
        except PermissionError:
            raise ExcelAbierto(
                "No se pudo guardar: el archivo vetgestion.xlsx está abierto en Excel. "
                "Cerralo y volvé a intentar.")

    def crear(self, prods, movs):
        wb = Workbook()
        ws = wb.active
        ws.title = "Productos"
        _estilar(ws, COLS_PROD)
        for p in prods:
            ws.append(_a_fila(p, COLS_PROD))
        wm = wb.create_sheet("Movimientos")
        _estilar(wm, COLS_MOV)
        for m in movs:
            wm.append(self._mov_excel(m))
        self._guardar(wb)

    def leer(self):
        wb = load_workbook(XLSX)
        return _filas_excel(wb["Productos"], COLS_PROD), _filas_excel(wb["Movimientos"], COLS_MOV)

    def aplicar(self, upserts, borrar, movs, parciales=()):
        upserts = list(upserts) + [c["fila"] for c in parciales]  # en Excel se guarda la fila completa
        wb = load_workbook(XLSX)
        ws = wb["Productos"]
        for p in upserts:
            fila = _a_fila(p, COLS_PROD)
            r = _buscar(ws, p["codigo"])
            if r:
                for i, v in enumerate(fila, start=1):
                    ws.cell(row=r, column=i, value=v)
            else:
                ws.append(fila)
        for codigo in borrar:
            r = _buscar(ws, codigo)
            if r:
                ws.delete_rows(r)
        for m in movs:
            wb["Movimientos"].append(self._mov_excel(m))
        self._guardar(wb)


# ---------------------------------------------------------------- almacén: Google Sheets
class SheetsStore:
    """Habla con el script de Apps Script publicado en la planilla (ver ui/apps_script.gs)."""
    modo = "sheets"
    nombre = "Google Sheets"

    def __init__(self, url):
        self.url = url.strip()
        self.version = 1  # la versión 2 del script aplica cambios puntuales y devuelve los datos (más rápido)

    def _post(self, accion, timeout=45, **datos):
        cuerpo = json.dumps({"accion": accion, **datos}, ensure_ascii=False).encode("utf-8")
        req = urllib.request.Request(self.url, data=cuerpo, headers={"Content-Type": "text/plain; charset=utf-8"})
        try:
            with urllib.request.urlopen(req, timeout=timeout) as r:
                texto = r.read().decode("utf-8")
        except urllib.error.HTTPError as e:
            raise ValueError(f"Google Sheets respondió con un error ({e.code}). Revisá que la URL sea la de la "
                             "aplicación web y que el acceso sea «Cualquier persona».")
        except (urllib.error.URLError, TimeoutError, OSError):
            raise SinConexion("No hay conexión con Google Sheets. Revisá internet e intentá de nuevo.")
        try:
            res = json.loads(texto)
        except ValueError:
            raise ValueError("La URL no responde como el script de VetGestión. Revisá que sea la URL de la "
                             "aplicación web (termina en /exec) y que el acceso sea «Cualquier persona».")
        if not res.get("ok"):
            if res.get("usuario"):  # p. ej. "No alcanza el stock": se muestra tal cual
                raise ValueError(str(res.get("error")))
            raise ValueError("Error en Google Sheets: " + str(res.get("error", "desconocido")))
        if "version" in res:
            self.version = int(res["version"])
        return res

    def ping(self):
        return self._post("ping", timeout=30)

    def inicializar(self, prods, movs):
        def hoja(cols, filas):
            return {"encabezados": [t for _, t, _ in cols],
                    "formatos": ['"$"#,##0' if k in CON_PESOS else "0" if k in NUMERICAS else "@" for k, _, _ in cols],
                    "anchos": [a * 7 for _, _, a in cols],
                    "filas": [_a_fila(x, cols) for x in filas]}
        return self._post("inicializar", timeout=90, prod=hoja(COLS_PROD, prods), mov=hoja(COLS_MOV, movs))

    @staticmethod
    def _datos(r):
        pk, mk = _claves(COLS_PROD), _claves(COLS_MOV)
        prods = [dict(zip(pk, f)) for f in r.get("productos", []) if f and str(f[0]).strip()]
        movs = [dict(zip(mk, f)) for f in r.get("movimientos", []) if f and any(str(x).strip() for x in f)]
        return prods, movs

    def leer(self):
        return self._datos(self._post("leer"))

    def aplicar(self, upserts, borrar, movs, parciales=()):
        """Con el script v2 devuelve los datos actualizados (prods, movs); con el v1, None."""
        filas = lambda xs: [_a_fila(p, COLS_PROD) for p in xs]
        comun = dict(borrar=list(borrar), movs=[_a_fila(m, COLS_MOV) for m in movs])
        if self.version >= 2:
            claves = _claves(COLS_PROD)
            r = self._post("aplicar", upserts=filas(upserts), devolver=True,
                           colStock=claves.index("stock") + 1, colPrecio=claves.index("precio") + 1,
                           parciales=[{"codigo": c["codigo"], "stock": c.get("stock"), "precio": c.get("precio")}
                                      for c in parciales], **comun)
            return self._datos(r)
        self._post("aplicar", upserts=filas(list(upserts) + [c["fila"] for c in parciales]), **comun)
        return None

    def subir_imagen(self, nombre, mime, datos):
        return self._post("subir_imagen", timeout=90, nombre=nombre, mime=mime,
                          datos=base64.b64encode(datos).decode())["id"]

    def imagen(self, id_drive):
        r = self._post("imagen", id=id_drive)
        return r["mime"], base64.b64decode(r["datos"])


# ---------------------------------------------------------------- configuración
def _leer_config():
    for ruta in (CONFIG, os.path.join(_recursos_dir(), "config.json")):
        if os.path.isfile(ruta):
            try:
                with open(ruta, encoding="utf-8") as f:
                    return json.load(f)
            except (OSError, ValueError):
                pass
    return {}


def _guardar_config(cfg):
    os.makedirs(DATOS_DIR, exist_ok=True)
    with open(CONFIG, "w", encoding="utf-8") as f:
        json.dump(cfg, f, ensure_ascii=False, indent=2)


_store = None
_config = {}
_cache = {"t": 0.0, "prods": None, "movs": None}


def store():
    return _store


def info_almacen():
    return {"modo": _store.modo, "nombre": _store.nombre,
            "link": _config.get("sheets_link", "") if _store.modo == "sheets" else XLSX,
            "planilla": _config.get("sheets_nombre", "")}


def _elegir_store():
    global _store, _config
    _config = _leer_config()
    url = _config.get("sheets_url", "")
    _store = SheetsStore(url) if url else ExcelStore()
    _cache["t"] = 0.0


def inicializar():
    """Crea la carpeta de datos y el Excel de respaldo, y elige dónde se guardan los datos."""
    viejo = os.path.join(_base_dir(), "datos")  # versiones anteriores guardaban todo junto al .exe
    if getattr(sys, "frozen", False) and os.path.isdir(viejo) and not os.path.exists(XLSX):
        shutil.copytree(viejo, DATOS_DIR, dirs_exist_ok=True)
    os.makedirs(IMG_DIR, exist_ok=True)
    src_img = os.path.join(_recursos_dir(), "ui", "img")
    if os.path.isdir(src_img):
        for f in os.listdir(src_img):
            dst = os.path.join(IMG_DIR, f)
            if not os.path.exists(dst):
                shutil.copyfile(os.path.join(src_img, f), dst)
    if not os.path.exists(XLSX):
        # el historial arranca vacío: solo se registra lo que pase de verdad en el programa
        ExcelStore().crear([_limpiar_producto(dict(zip(_claves(COLS_PROD), f))) for f in SEED], [])
    _elegir_store()


def _es_semilla(nombre):
    return os.path.isfile(os.path.join(_recursos_dir(), "ui", "img", os.path.basename(nombre)))


def conectar_sheets(url):
    """Pasa a trabajar con Google Sheets. Si la planilla está vacía, le sube los datos actuales."""
    url = url.strip()
    if not url.startswith("https://script.google.com/") or not url.rstrip("/").endswith("/exec"):
        raise ValueError("La URL tiene que ser la de la aplicación web de Apps Script: empieza con "
                         "https://script.google.com/ y termina en /exec.")
    nuevo = SheetsStore(url.rstrip("/"))
    info = nuevo.ping()
    with _lock:
        subidos = False
        if not info.get("productos"):
            prods, movs = _leer(fresco=True)
            for p in prods:  # las fotos cargadas a mano se suben a Drive para verlas desde cualquier PC
                img = p["imagen"]
                ruta = os.path.join(IMG_DIR, os.path.basename(img))
                if img and not img.startswith("drive:") and not _es_semilla(img) and os.path.isfile(ruta):
                    with open(ruta, "rb") as f:
                        mime = "image/" + {"jpg": "jpeg"}.get(img.rsplit(".", 1)[-1].lower(), img.rsplit(".", 1)[-1].lower())
                        p["imagen"] = "drive:" + nuevo.subir_imagen(img, mime, f.read())
            nuevo.inicializar(prods, movs)
            subidos = True
        else:
            nuevo.inicializar([], [])  # asegura encabezados y formato, sin tocar los datos
        _config.update(sheets_url=nuevo.url, sheets_link=info.get("url", ""), sheets_nombre=info.get("nombre", ""))
        _guardar_config(_config)
        _elegir_store()
    return {"subidos": subidos, "planilla": info.get("nombre", "")}


def desconectar_sheets():
    """Vuelve al Excel local, copiando antes los datos de la nube para no perder nada."""
    with _lock:
        copiados = False
        if _store.modo == "sheets":
            try:
                prods, movs = _leer(fresco=True)
                for p in prods:
                    if p["imagen"].startswith("drive:"):
                        p["imagen"] = imagen_local(p["imagen"], _store) or ""
                ExcelStore().crear(prods, movs)
                copiados = True
            except SinConexion:
                pass
        _guardar_config({"sheets_url": ""})
        _elegir_store()
    return {"copiados": copiados}


def imagen_local(nombre, almacen=None):
    """Ruta local de una imagen. Las de Drive se descargan una vez y quedan guardadas."""
    if not nombre.startswith("drive:"):
        ruta = os.path.join(IMG_DIR, os.path.basename(nombre))
        return os.path.basename(ruta) if os.path.isfile(ruta) else None
    id_drive = nombre[6:]
    for f in os.listdir(IMG_DIR):
        if f.startswith(f"drive_{id_drive}."):
            return f
    almacen = almacen or _store
    if almacen.modo != "sheets":
        return None
    mime, datos = almacen.imagen(id_drive)
    ext = {"image/jpeg": "jpg", "image/png": "png", "image/webp": "webp", "image/gif": "gif"}.get(mime, "jpg")
    local = f"drive_{id_drive}.{ext}"
    with open(os.path.join(IMG_DIR, local), "wb") as f:
        f.write(datos)
    return local


# ---------------------------------------------------------------- lectura / escritura con caché
COPIA = os.path.join(DATOS_DIR, "copia_nube.json")


def _guardar_copia():
    """Copia local de lo último leído de la nube: permite abrir el programa al instante."""
    if _store.modo != "sheets":
        return
    try:
        with open(COPIA, "w", encoding="utf-8") as f:
            json.dump({"url": _store.url, "prods": _cache["prods"], "movs": _cache["movs"]}, f, ensure_ascii=False)
    except OSError:
        pass


def leer_copia(limite_mov=None):
    """Datos de la última lectura de la nube (pueden tener algunos minutos). None si no hay copia."""
    if _store.modo != "sheets" or not os.path.isfile(COPIA):
        return None
    try:
        with open(COPIA, encoding="utf-8") as f:
            c = json.load(f)
    except (OSError, ValueError):
        return None
    if c.get("url") != _store.url:
        return None
    return _preparar([_limpiar_producto(p) for p in c["prods"]], [_limpiar_mov(m) for m in c["movs"]], limite_mov)


def _leer(fresco=False):
    if fresco or _cache["prods"] is None or time.time() - _cache["t"] > CACHE_SEG:
        prods, movs = _store.leer()
        _cache.update(t=time.time(), prods=[_limpiar_producto(p) for p in prods],
                      movs=[_limpiar_mov(m) for m in movs])
        _guardar_copia()
    return copy.deepcopy(_cache["prods"]), copy.deepcopy(_cache["movs"])


def _leer_para_escribir():
    """Con el script v2 alcanza con los datos en memoria (se actualizan tras cada operación y cada
    30 s): los cambios de stock y precio se aplican sobre lo que hay en la planilla en ese momento."""
    rapido = _store.modo == "sheets" and _store.version >= 2 and _cache["prods"] is not None \
        and time.time() - _cache["t"] < 90
    return _leer(fresco=not rapido)


def _escribir(upserts=(), borrar=(), movs=(), parciales=()):
    """parciales: [{codigo, stock (delta) y/o precio (nuevo), fila (producto completo ya modificado)}]."""
    upserts, borrar, movs, parciales = list(upserts), list(borrar), list(movs), list(parciales)
    try:
        nuevos = _store.aplicar(upserts, borrar, movs, parciales)
    except Exception:
        _cache["t"] = 0.0  # ante cualquier error, la próxima lectura va a la fuente
        raise
    if nuevos is not None:  # el script devolvió los datos actualizados
        prods, mvs = nuevos
        _cache.update(t=time.time(), prods=[_limpiar_producto(p) for p in prods],
                      movs=[_limpiar_mov(m) for m in mvs])
        _guardar_copia()
        return
    upserts += [c["fila"] for c in parciales]
    prods = _cache["prods"]
    for p in upserts:
        i = next((n for n, x in enumerate(prods) if x["codigo"] == p["codigo"]), None)
        if i is None:
            prods.append(copy.deepcopy(p))
        else:
            prods[i] = copy.deepcopy(p)
    _cache["prods"] = [p for p in prods if p["codigo"] not in set(borrar)]
    _cache["movs"].extend(copy.deepcopy(movs))
    _cache["t"] = time.time()
    _guardar_copia()


def _mov(tipo, codigo="", producto="", cantidad=0, importe=0, detalle="", usuario=""):
    return {"fecha": datetime.now().strftime(FMT_FECHA), "tipo": tipo, "codigo": codigo, "producto": producto,
            "cantidad": cantidad, "importe": importe, "detalle": detalle, "usuario": usuario}


def _producto(prods, codigo):
    p = next((x for x in prods if x["codigo"] == codigo), None)
    if not p:
        raise ValueError("El producto ya no existe.")
    return p


def leer_todo(limite_mov=40):
    with _lock:
        prods, movs = _leer()
    return _preparar(prods, movs, limite_mov)


def _preparar(prods, movs, limite_mov):
    for p in prods:
        p["estado"], p["dias"] = estado_vencimiento(p["vencimiento"])
    movs.sort(key=lambda m: m["fecha"], reverse=True)
    for m in movs:
        m["fecha"] = m["fecha"].replace(" ", "T")
    return prods, (movs if limite_mov is None else movs[:limite_mov])


def vender(codigo, cantidad, usuario):
    with _lock:
        prods, _ = _leer_para_escribir()
        p = _producto(prods, codigo)
        if estado_vencimiento(p["vencimiento"])[0] == "vencido":
            raise ValueError("Este producto está vencido: no se puede vender.")
        if cantidad < 1:
            raise ValueError("La cantidad tiene que ser al menos 1.")
        if cantidad > p["stock"]:
            raise ValueError(f"No alcanza el stock: quedan {p['stock']} unidades.")
        p["stock"] -= cantidad
        _escribir(parciales=[{"codigo": codigo, "stock": -cantidad, "fila": p}],
                  movs=[_mov("Venta", codigo, p["nombre"], cantidad, p["precio"] * cantidad, "", usuario)])
        quedan = next((x["stock"] for x in _cache["prods"] if x["codigo"] == codigo), p["stock"])
        return {"nombre": p["nombre"], "quedan": quedan, "total": p["precio"] * cantidad}


def ajustar_stock(codigo, cantidad, tipo, detalle, usuario):
    """Suma (reposición) o resta (retiro) unidades de stock."""
    with _lock:
        prods, _ = _leer_para_escribir()
        p = _producto(prods, codigo)
        if p["stock"] + cantidad < 0:
            raise ValueError(f"No se pueden retirar {-cantidad}: hay {p['stock']} unidades.")
        p["stock"] += cantidad
        _escribir(parciales=[{"codigo": codigo, "stock": cantidad, "fila": p}],
                  movs=[_mov(tipo, codigo, p["nombre"], abs(cantidad), 0, detalle, usuario)])
        stock = next((x["stock"] for x in _cache["prods"] if x["codigo"] == codigo), p["stock"])
        return {"nombre": p["nombre"], "stock": stock}


def aumentar(reglas, usuario):
    """reglas: lista de (laboratorio, categoría o '', porcentaje). Devuelve resumen."""
    with _lock:
        prods, _ = _leer_para_escribir()
        cambiados, movs, resumen = {}, [], []
        for lab, cat, pct in reglas:
            n = total = 0
            for p in prods:
                if p["laboratorio"].lower() != lab.strip().lower() or (cat and p["categoria"] != cat):
                    continue
                p["precio"] = round(p["precio"] * (1 + pct / 100))
                cambiados[p["codigo"]] = p
                n += 1
                total += p["precio"]
            if n:
                pct_txt = f"{pct:g}".replace(".", ",")
                movs.append(_mov("Aumento de precio", "", lab, n, 0,
                                 f"{lab} {'+' if pct > 0 else ''}{pct_txt} %" + (f" · {cat}" if cat else ""), usuario))
            resumen.append({"laboratorio": lab, "categoria": cat, "pct": pct, "productos": n, "total": total})
        if not cambiados:
            raise ValueError("Ningún producto coincide con esos laboratorios.")
        _escribir(parciales=[{"codigo": c, "precio": p["precio"], "fila": p} for c, p in cambiados.items()], movs=movs)
        return resumen


def guardar_producto(datos, usuario, codigo_original=None):
    """Alta (codigo_original=None) o edición de un producto."""
    with _lock:
        prods, _ = _leer(fresco=True)
        codigo = datos["codigo"]
        existe = any(p["codigo"] == codigo for p in prods)
        if codigo_original:
            anterior = _producto(prods, codigo_original)
            if codigo != codigo_original and existe:
                raise ValueError(f"Ya existe otro producto con el código {codigo}.")
            nuevo = _limpiar_producto({**anterior, **datos})
            cambios = []
            if anterior["precio"] != nuevo["precio"]:
                cambios.append(f"precio ${anterior['precio']:,} → ${nuevo['precio']:,}".replace(",", "."))
            if anterior["stock"] != nuevo["stock"]:
                cambios.append(f"stock {anterior['stock']} → {nuevo['stock']}")
            _escribir([nuevo], borrar=[codigo_original] if codigo != codigo_original else [],
                      movs=[_mov("Edición de producto", codigo, nuevo["nombre"], 0, 0,
                                 ", ".join(cambios) or "Datos de la ficha", usuario)])
        else:
            if existe:
                raise ValueError(f"Ya existe un producto con el código {codigo}.")
            nuevo = _limpiar_producto(datos)
            _escribir([nuevo], movs=[_mov("Alta de producto", codigo, nuevo["nombre"], nuevo["stock"], 0,
                                          f"{nuevo['laboratorio']} · {nuevo['categoria']}", usuario)])


def baja(codigo, usuario):
    with _lock:
        prods, _ = _leer(fresco=True)
        p = _producto(prods, codigo)
        _escribir(borrar=[codigo], movs=[_mov("Baja de producto", codigo, p["nombre"], p["stock"], 0,
                                              "Eliminado del catálogo", usuario)])
        return p["nombre"]
